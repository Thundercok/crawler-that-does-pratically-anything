"""
rat.crawler.extractors — Deep content extraction for diverse document formats.
"""

from __future__ import annotations

import csv
import io
import logging
import os
from pathlib import Path
from typing import Dict, Any, Optional

logger = logging.getLogger("rat.extractors")

MAX_CHARS_PER_DOC = 100_000  # Cap content length to keep search snappy


def extract_text_from_pdf(file_path: str) -> str:
    """Extract text and metadata from a PDF file using pypdf."""
    try:
        from pypdf import PdfReader
        reader = PdfReader(file_path)
        texts = []
        # Extract metadata
        meta = reader.metadata
        if meta:
            if meta.title:
                texts.append(f"Title: {meta.title}")
            if meta.author:
                texts.append(f"Author: {meta.author}")
            if meta.subject:
                texts.append(f"Subject: {meta.subject}")

        # Extract text per page
        for idx, page in enumerate(reader.pages):
            page_text = page.extract_text()
            if page_text and page_text.strip():
                texts.append(f"--- Page {idx + 1} ---\n{page_text.strip()}")
            if sum(len(t) for t in texts) > MAX_CHARS_PER_DOC:
                break
        return "\n\n".join(texts)
    except Exception as e:
        logger.debug(f"Failed to extract PDF {file_path}: {e}")
        return ""


def extract_text_from_docx(file_path: str) -> str:
    """Extract text from Word (.docx) file using python-docx."""
    try:
        import docx
        doc = docx.Document(file_path)
        texts = []

        # Extract core properties if available
        try:
            core_props = doc.core_properties
            if core_props.title:
                texts.append(f"Title: {core_props.title}")
            if core_props.author:
                texts.append(f"Author: {core_props.author}")
            if core_props.comments:
                texts.append(f"Comments: {core_props.comments}")
        except Exception:
            pass

        # Extract paragraphs
        for p in doc.paragraphs:
            if p.text and p.text.strip():
                texts.append(p.text.strip())
            if sum(len(t) for t in texts) > MAX_CHARS_PER_DOC:
                break

        # Extract tables
        for table in doc.tables:
            for row in table.rows:
                row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_texts:
                    texts.append(" | ".join(row_texts))
            if sum(len(t) for t in texts) > MAX_CHARS_PER_DOC:
                break

        return "\n".join(texts)
    except Exception as e:
        logger.debug(f"Failed to extract DOCX {file_path}: {e}")
        return ""


def extract_text_from_pptx(file_path: str) -> str:
    """Extract text from PowerPoint (.pptx) file using python-pptx."""
    try:
        from pptx import Presentation
        prs = Presentation(file_path)
        texts = []
        for idx, slide in enumerate(prs.slides):
            slide_texts = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    slide_texts.append(shape.text.strip())
            if slide_texts:
                texts.append(f"--- Slide {idx + 1} ---\n" + "\n".join(slide_texts))
            if sum(len(t) for t in texts) > MAX_CHARS_PER_DOC:
                break
        return "\n\n".join(texts)
    except Exception as e:
        logger.debug(f"Failed to extract PPTX {file_path}: {e}")
        return ""


def extract_text_from_xlsx(file_path: str) -> str:
    """Extract sheet names and sample cell values from Excel (.xlsx) file."""
    try:
        import openpyxl
        wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
        texts = [f"Sheets: {', '.join(wb.sheetnames)}"]
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            texts.append(f"--- Sheet: {sheet_name} ---")
            row_count = 0
            for row in ws.iter_rows(values_only=True):
                row_vals = [str(v).strip() for v in row if v is not None and str(v).strip()]
                if row_vals:
                    texts.append(" | ".join(row_vals))
                row_count += 1
                if row_count > 200 or sum(len(t) for t in texts) > MAX_CHARS_PER_DOC:
                    break
        wb.close()
        return "\n".join(texts)
    except Exception as e:
        logger.debug(f"Failed to extract XLSX {file_path}: {e}")
        return ""


def extract_text_from_csv(file_path: str) -> str:
    """Extract rows from a CSV file."""
    try:
        texts = []
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            reader = csv.reader(f)
            for idx, row in enumerate(reader):
                if row:
                    texts.append(" | ".join(row))
                if idx > 200:
                    break
        return "\n".join(texts)
    except Exception as e:
        logger.debug(f"Failed to extract CSV {file_path}: {e}")
        return ""


def extract_text_from_plaintext(file_path: str) -> str:
    """Extract text from code, markdown, txt, json, etc."""
    encodings = ["utf-8", "utf-8-sig", "latin-1", "cp1252", "cp1258"]
    for enc in encodings:
        try:
            with open(file_path, "r", encoding=enc) as f:
                content = f.read(MAX_CHARS_PER_DOC)
                return content
        except Exception:
            continue
    return ""


def extract_text_from_image(file_path: str) -> str:
    """
    Extract multimodal image content: Native Apple Vision OCR + Scene/Object Classification + EXIF.
    """
    texts = []
    try:
        from rat.crawler.apple_vision import apple_vision
        if apple_vision.available:
            analysis = apple_vision.analyze_image(file_path)
            if analysis.get("combined_text"):
                return analysis["combined_text"]

        # Fallback for non-macOS or if Vision fails
        from PIL import Image
        size_bytes = os.path.getsize(file_path)
        if size_bytes > 15 * 1024 * 1024:
            return f"Image file: {Path(file_path).name}, Size: {size_bytes} bytes"

        with Image.open(file_path) as img:
            texts.append(f"Image format: {img.format}, Size: {img.width}x{img.height}")
            try:
                exif = img._getexif()
                if exif:
                    for tag_id, val in exif.items():
                        if isinstance(val, (str, int, float)) and len(str(val)) < 100:
                            texts.append(f"EXIF {tag_id}: {val}")
            except Exception:
                pass

            if img.width <= 4000 and img.height <= 4000:
                try:
                    import pytesseract
                    ocr_text = pytesseract.image_to_string(img, timeout=2)
                    if ocr_text.strip():
                        texts.append("OCR Content:\n" + ocr_text.strip())
                except Exception:
                    pass
    except Exception as e:
        logger.debug(f"Failed to read image {file_path}: {e}")

    return "\n".join(texts)


def extract_document_content(file_path: str) -> str:
    """
    Dispatcher to extract deep text content based on file extension.
    """
    ext = Path(file_path).suffix.lower()

    if ext == ".pdf":
        return extract_text_from_pdf(file_path)
    elif ext in [".docx", ".doc"]:
        return extract_text_from_docx(file_path)
    elif ext in [".pptx", ".ppt"]:
        return extract_text_from_pptx(file_path)
    elif ext in [".xlsx", ".xls"]:
        return extract_text_from_xlsx(file_path)
    elif ext == ".csv":
        return extract_text_from_csv(file_path)
    elif ext in [".jpg", ".jpeg", ".png", ".webp"]:
        return extract_text_from_image(file_path)
    elif ext in [
        ".txt", ".md", ".markdown", ".rst", ".py", ".js", ".jsx", ".ts", ".tsx",
        ".html", ".css", ".json", ".yaml", ".yml", ".toml", ".sh", ".sql",
        ".xml", ".log", ".rtf"
    ]:
        return extract_text_from_plaintext(file_path)
    else:
        return ""
